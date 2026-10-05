#!/usr/bin/env python3
"""
RERA FORENSIC & COMPLIANCE AUDIT INTELLIGENCE SYSTEM
====================================================
Canonical Data Ingestion & Normalization Engine
Converts Schema A (Verified) and Schema B (Compared) into Canonical RERA Schema.

Supported Projects:
1. DREAM COLONY (Schema A)
2. RIDHI SIDHI COMMERCIAL HUB BLOCK A (Schema B)
3. SHIV ENCLAVE (Schema A)

Author: Principal Forensic Data Engineer
"""

import os
import re
import csv
import json
import glob
from datetime import datetime
import openpyxl

QUARTER_ORDER = {
    'January to March': 1,
    'April to June': 2,
    'July to September': 3,
    'October to December': 4
}

def get_quarter_sort_key(quarter_str: str) -> int:
    parts = quarter_str.split(',')
    if len(parts) == 2:
        period = parts[0].strip()
        year = int(parts[1].strip())
        q_num = QUARTER_ORDER.get(period, 0)
        return year * 10 + q_num
    return 0

def clean_currency(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s in ('-', 'N/A', 'None', 'null', 'Not Available', 'ignored', 'on', 'Work in Progress'):
        return None
    # Strip ₹, commas, spaces
    s = re.sub(r'[₹,\s]', '', s)
    # Strip date concatenation artifacts
    s = re.sub(r'^\d{2}-\d{2}-\d{4}', '', s)
    try:
        return float(s)
    except ValueError:
        return None

def clean_percentage(val) -> float | None:
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip()
    if not s or s in ('-', 'N/A', 'None', 'null', 'Not Available', 'Not Yet Started', 'Work in Progress'):
        return 0.0
    s = re.sub(r'[%\s,]', '', s)
    try:
        return float(s)
    except ValueError:
        return None

def calculate_severity(metric: str, web_val: float | None, pdf_val: float | None, unit: str, is_missing_doc=False, is_ignored=False):
    if is_ignored:
        return 'Ignored', 'IGNORED', 'Standard portal placeholder row excluded from audit statistics'
    
    if is_missing_doc:
        return 'Missing Document', 'DATA GAP', 'Required certified legal evidence was not found or was not uploaded to portal'

    if web_val is None or pdf_val is None:
        return 'Unverified', 'WARNING', 'Data point is missing in either the public portal or the certified document'

    abs_var = abs(pdf_val - web_val)

    if unit == '%':
        if abs_var < 0.5:
            return 'Match', 'MATCH', 'Public declaration matches certified evidence within 0.5% tolerance'
        elif abs_var > 30.0:
            return 'Mismatch', 'CRITICAL', f'Physical progress discrepancy of {abs_var:.1f} percentage points exceeds the 30 pp critical threshold'
        else:
            return 'Mismatch', 'WARNING', f'Physical progress discrepancy of {abs_var:.1f} percentage points detected'
    else:
        # Currency (₹)
        if abs_var <= 100:
            return 'Match', 'MATCH', 'Public declaration matches certified evidence within ₹100 rounding tolerance'
        elif abs_var >= 1_000_000:
            cr_val = abs_var / 10_000_000
            return 'Mismatch', 'CRITICAL', f'Financial discrepancy of ₹{cr_val:.2f} Cr exceeds the ₹10 Lakh critical threshold'
        else:
            lakh_val = abs_var / 100_000
            return 'Mismatch', 'WARNING', f'Financial discrepancy of ₹{lakh_val:.2f} Lakh detected below critical threshold'

def parse_inventory(csv_path: str) -> dict:
    inv = {'total': 0, 'sold': 0, 'unsold': 0, 'mortgaged': 0}
    if not os.path.exists(csv_path):
        return inv

    with open(csv_path, 'r', encoding='utf-8', errors='ignore') as f:
        reader = csv.DictReader(f)
        for row in reader:
            status = row.get('Booking Status', '').strip().lower()
            inv['total'] += 1
            if 'sold' in status and 'unsold' not in status:
                inv['sold'] += 1
            elif 'unsold' in status:
                inv['unsold'] += 1
            elif 'mortgage' in status:
                inv['mortgaged'] += 1
            else:
                inv['unsold'] += 1
    return inv

def find_file(q_dir: str, candidates: list[str]) -> str | None:
    for c in candidates:
        p = os.path.join(q_dir, c)
        if os.path.exists(p):
            return p
    return None

def process_project(project_folder: str, project_id: str, project_name: str, promoter: str, district: str, project_type: str, reg_no: str, project_url: str):
    sub_folder = os.path.join(project_folder, 'sub')
    sub_sub = os.listdir(sub_folder)[0]
    base_path = os.path.join(sub_folder, sub_sub)

    quarter_dirs = [d for d in os.listdir(base_path) if os.path.isdir(os.path.join(base_path, d))]

    quarters_data = {}
    project_findings = []
    max_estimated_capital = 0.0

    for q_name in quarter_dirs:
        q_dir = os.path.join(base_path, q_name)
        sort_key = get_quarter_sort_key(q_name)

        # Check for inventory
        csv_file = os.path.join(q_dir, 'Allottee_Details.csv')
        inventory = parse_inventory(csv_file)

        # Candidate files
        r1_file = find_file(q_dir, ['Form_R1_Verified.xlsx', 'Form_R1_Compared.xlsx', 'Form_R1.xlsx'])
        r2_file = find_file(q_dir, ['Form_R2_Verified.xlsx', 'Form_R2_Compared.xlsx', 'Form_R2.xlsx'])
        r3_top_file = find_file(q_dir, ['Form_R3_TOP_Verified.xlsx', 'Form_R3_TOP_Compared.xlsx', 'Form_R3_TOP.xlsx', 'Form_R3_Top.xlsx'])
        r3_bot_file = find_file(q_dir, ['Form_R3_BOTTOM_Verified.xlsx', 'Form_R3_BOTTOM_Compared.xlsx', 'Form_R3_BOTTOM.xlsx', 'Form_R3_Bottom.xlsx'])

        is_active = (r1_file is not None or r2_file is not None or r3_top_file is not None or r3_bot_file is not None)

        form_r1_records = []
        form_r2_records = []
        form_r3_top_records = []
        form_r3_bottom_records = []

        q_stats = {
            'matches': 0,
            'mismatches': 0,
            'critical': 0,
            'warnings': 0,
            'data_gaps': 0,
            'ignored': 0
        }

        # --- FORM R-1 ---
        if r1_file:
            wb = openpyxl.load_workbook(r1_file, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows and len(rows) > 1:
                header = [str(c).strip() if c else '' for c in rows[0]]
                task_idx = header.index('Task/Facility') if 'Task/Facility' in header else 2
                web_idx = header.index('Percentage (Current Quarter)') if 'Percentage (Current Quarter)' in header else 4
                
                pdf_idx = None
                if 'Extracted_PDF_Percentage' in header:
                    pdf_idx = header.index('Extracted_PDF_Percentage')
                elif 'PDF_Percentage' in header:
                    pdf_idx = header.index('PDF_Percentage')

                doc_idx = header.index('docs') if 'docs' in header else -1

                for row_idx, r in enumerate(rows[1:], start=1):
                    task = str(r[task_idx] or '').strip()
                    if not task:
                        continue
                    
                    web_raw = str(r[web_idx]) if web_idx < len(r) and r[web_idx] is not None else ''
                    pdf_raw = str(r[pdf_idx]) if pdf_idx is not None and pdf_idx < len(r) and r[pdf_idx] is not None else ''
                    doc_url = str(r[doc_idx]) if doc_idx != -1 and doc_idx < len(r) and r[doc_idx] is not None else ''

                    web_val = clean_percentage(web_raw)
                    pdf_val = clean_percentage(pdf_raw) if (pdf_raw and pdf_raw != 'None') else None

                    is_missing_doc = (not doc_url or 'not found' in doc_url.lower() or not doc_url.startswith('http') or pdf_val is None and not pdf_raw)
                    if 'not found' in doc_url.lower() or not doc_url or not doc_url.startswith('http'):
                        is_missing_doc = True

                    status, severity, reason = calculate_severity(task, web_val, pdf_val, '%', is_missing_doc=is_missing_doc)

                    var_abs = (pdf_val - web_val) if (pdf_val is not None and web_val is not None) else None
                    var_rel = (abs(pdf_val - web_val) / max(abs(pdf_val), abs(web_val), 1.0) * 100) if (pdf_val is not None and web_val is not None) else None

                    rec = {
                        'id': f"R1-{project_id}-{sort_key}-{row_idx}",
                        'task': task,
                        'web_raw': web_raw,
                        'pdf_raw': pdf_raw,
                        'web_val': web_val,
                        'pdf_val': pdf_val,
                        'variance': var_abs,
                        'variance_pp': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'docs_url': doc_url
                    }
                    form_r1_records.append(rec)

                    if status == 'Match':
                        q_stats['matches'] += 1
                    elif status == 'Mismatch':
                        q_stats['mismatches'] += 1
                    elif status == 'Missing Document':
                        q_stats['data_gaps'] += 1

                    if severity == 'CRITICAL':
                        q_stats['critical'] += 1
                    elif severity == 'WARNING':
                        q_stats['warnings'] += 1

                    project_findings.append({
                        'id': f"FND-R1-{project_id}-{sort_key}-{row_idx}",
                        'project_id': project_id,
                        'project_name': project_name,
                        'district': district,
                        'promoter': promoter,
                        'quarter': q_name,
                        'sort_key': sort_key,
                        'form_type': 'Form R-1 (Architect)',
                        'metric': task,
                        'category': 'Physical Infrastructure',
                        'unit': '%',
                        'web_raw': web_raw,
                        'pdf_raw': pdf_raw,
                        'web_val': web_val,
                        'pdf_val': pdf_val,
                        'absolute_variance': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'document_url': doc_url,
                        'document_status': 'Missing' if is_missing_doc else 'Available'
                    })

        # --- FORM R-2 ---
        if r2_file:
            wb = openpyxl.load_workbook(r2_file, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows and len(rows) > 1:
                header = [str(c).strip() if c else '' for c in rows[0]]
                part_idx = header.index('Particulars') if 'Particulars' in header else 1
                web_idx = header.index('In Current Quarter') if 'In Current Quarter' in header else 3
                
                pdf_idx = None
                if 'Extracted_PDF_Percentage' in header:
                    pdf_idx = header.index('Extracted_PDF_Percentage')
                elif 'PDF_Percentage' in header:
                    pdf_idx = header.index('PDF_Percentage')
                elif 'PDF_Extracted_Amount' in header:
                    pdf_idx = header.index('PDF_Extracted_Amount')
                elif 'Extracted_PDF_Amount' in header:
                    pdf_idx = header.index('Extracted_PDF_Amount')

                doc_idx = header.index('docs') if 'docs' in header else -1

                for row_idx, r in enumerate(rows[1:], start=1):
                    part = str(r[part_idx] or '').strip()
                    if not part:
                        continue
                    
                    unit = '%' if 'Percentage' in part or '%' in part else '₹'
                    web_raw = str(r[web_idx]) if web_idx < len(r) and r[web_idx] is not None else ''
                    pdf_raw = str(r[pdf_idx]) if pdf_idx is not None and pdf_idx < len(r) and r[pdf_idx] is not None else ''
                    doc_url = str(r[doc_idx]) if doc_idx != -1 and doc_idx < len(r) and r[doc_idx] is not None else ''

                    web_val = clean_percentage(web_raw) if unit == '%' else clean_currency(web_raw)
                    pdf_val = (clean_percentage(pdf_raw) if unit == '%' else clean_currency(pdf_raw)) if (pdf_raw and pdf_raw != 'None') else None

                    is_missing_doc = (not doc_url or 'not found' in doc_url.lower() or not doc_url.startswith('http') or pdf_val is None and not pdf_raw)
                    status, severity, reason = calculate_severity(part, web_val, pdf_val, unit, is_missing_doc=is_missing_doc)

                    var_abs = (pdf_val - web_val) if (pdf_val is not None and web_val is not None) else None
                    var_rel = (abs(pdf_val - web_val) / max(abs(pdf_val), abs(web_val), 1.0) * 100) if (pdf_val is not None and web_val is not None) else None

                    rec = {
                        'id': f"R2-{project_id}-{sort_key}-{row_idx}",
                        'particulars': part,
                        'unit': unit,
                        'web_raw': web_raw,
                        'pdf_raw': pdf_raw,
                        'web_val': web_val,
                        'pdf_val': pdf_val,
                        'variance': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'docs_url': doc_url
                    }
                    form_r2_records.append(rec)

                    if status == 'Match':
                        q_stats['matches'] += 1
                    elif status == 'Mismatch':
                        q_stats['mismatches'] += 1
                    elif status == 'Missing Document':
                        q_stats['data_gaps'] += 1

                    if severity == 'CRITICAL':
                        q_stats['critical'] += 1
                    elif severity == 'WARNING':
                        q_stats['warnings'] += 1

                    project_findings.append({
                        'id': f"FND-R2-{project_id}-{sort_key}-{row_idx}",
                        'project_id': project_id,
                        'project_name': project_name,
                        'district': district,
                        'promoter': promoter,
                        'quarter': q_name,
                        'sort_key': sort_key,
                        'form_type': 'Form R-2 (Engineer)',
                        'metric': part,
                        'category': 'Engineering Budget',
                        'unit': unit,
                        'web_raw': web_raw,
                        'pdf_raw': pdf_raw,
                        'web_val': web_val,
                        'pdf_val': pdf_val,
                        'absolute_variance': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'document_url': doc_url,
                        'document_status': 'Missing' if is_missing_doc else 'Available'
                    })

        # --- FORM R-3 TOP ---
        if r3_top_file:
            wb = openpyxl.load_workbook(r3_top_file, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows and len(rows) > 1:
                header = [str(c).strip() if c else '' for c in rows[0]]
                part_idx = header.index('Particulars') if 'Particulars' in header else 1
                est_idx = header.index('Estimated Cost (INR)') if 'Estimated Cost (INR)' in header else 2
                inc_web_idx = header.index('Incurred Cost (INR)') if 'Incurred Cost (INR)' in header else 3
                
                inc_pdf_idx = None
                if 'Extracted_PDF_Incurred' in header:
                    inc_pdf_idx = header.index('Extracted_PDF_Incurred')
                elif 'PDF_Extracted_Amount' in header:
                    inc_pdf_idx = header.index('PDF_Extracted_Amount')
                elif 'PDF_Incurred' in header:
                    inc_pdf_idx = header.index('PDF_Incurred')

                doc_idx = header.index('docs') if 'docs' in header else -1

                for row_idx, r in enumerate(rows[1:], start=1):
                    part = str(r[part_idx] or '').strip()
                    if not part:
                        continue
                    
                    web_est = clean_currency(r[est_idx]) if est_idx < len(r) else None
                    web_inc = clean_currency(r[inc_web_idx]) if inc_web_idx < len(r) else None
                    pdf_inc = clean_currency(r[inc_pdf_idx]) if (inc_pdf_idx is not None and inc_pdf_idx < len(r)) else None
                    doc_url = str(r[doc_idx]) if doc_idx != -1 and doc_idx < len(r) and r[doc_idx] is not None else ''

                    if web_est and ('Total' in part or 'Estimated Cost' in part):
                        max_estimated_capital = max(max_estimated_capital, web_est)

                    is_missing_doc = (not doc_url or 'not found' in doc_url.lower() or not doc_url.startswith('http') or (pdf_inc is None and (inc_pdf_idx is None or not r[inc_pdf_idx])))
                    status, severity, reason = calculate_severity(part, web_inc, pdf_inc, '₹', is_missing_doc=is_missing_doc)

                    var_abs = (pdf_inc - web_inc) if (pdf_inc is not None and web_inc is not None) else None
                    var_rel = (abs(pdf_inc - web_inc) / max(abs(pdf_inc), abs(web_inc), 1.0) * 100) if (pdf_inc is not None and web_inc is not None) else None

                    rec = {
                        'id': f"R3T-{project_id}-{sort_key}-{row_idx}",
                        'particulars': part,
                        'web_estimated': web_est,
                        'pdf_estimated': None,
                        'web_incurred': web_inc,
                        'pdf_incurred': pdf_inc,
                        'status_incurred': status,
                        'severity_incurred': severity,
                        'variance_incurred': abs(var_abs) if var_abs is not None else None,
                        'docs_url': doc_url
                    }
                    form_r3_top_records.append(rec)

                    if status == 'Match':
                        q_stats['matches'] += 1
                    elif status == 'Mismatch':
                        q_stats['mismatches'] += 1
                    elif status == 'Missing Document':
                        q_stats['data_gaps'] += 1

                    if severity == 'CRITICAL':
                        q_stats['critical'] += 1
                    elif severity == 'WARNING':
                        q_stats['warnings'] += 1

                    project_findings.append({
                        'id': f"FND-R3T-{project_id}-{sort_key}-{row_idx}",
                        'project_id': project_id,
                        'project_name': project_name,
                        'district': district,
                        'promoter': promoter,
                        'quarter': q_name,
                        'sort_key': sort_key,
                        'form_type': 'Form R-3 Table 1 (Cost Audit)',
                        'metric': part,
                        'category': 'Capital Cost Incurred',
                        'unit': '₹',
                        'web_raw': str(r[inc_web_idx]) if inc_web_idx < len(r) else '',
                        'pdf_raw': str(r[inc_pdf_idx]) if (inc_pdf_idx is not None and inc_pdf_idx < len(r)) else '',
                        'web_val': web_inc,
                        'pdf_val': pdf_inc,
                        'absolute_variance': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'document_url': doc_url,
                        'document_status': 'Missing' if is_missing_doc else 'Available'
                    })

        # --- FORM R-3 BOTTOM ---
        if r3_bot_file:
            wb = openpyxl.load_workbook(r3_bot_file, data_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows and len(rows) > 1:
                header = [str(c).strip() if c else '' for c in rows[0]]
                sno_idx = header.index('S.No.') if 'S.No.' in header else 0
                part_idx = header.index('Particulars') if 'Particulars' in header else 1
                unit_idx = header.index('Unit') if 'Unit' in header else 2
                web_idx = header.index('In Current Quarter') if 'In Current Quarter' in header else 4
                
                pdf_idx = None
                if 'PDF_Extracted_Amount' in header:
                    pdf_idx = header.index('PDF_Extracted_Amount')
                elif 'Extracted_PDF_Amount' in header:
                    pdf_idx = header.index('Extracted_PDF_Amount')
                elif 'PDF_Percentage' in header:
                    pdf_idx = header.index('PDF_Percentage')

                doc_idx = header.index('docs') if 'docs' in header else -1

                for row_idx, r in enumerate(rows[1:], start=1):
                    sno = str(r[sno_idx] if sno_idx < len(r) else row_idx).strip()
                    part = str(r[part_idx] or '').strip()
                    raw_unit = str(r[unit_idx] or '').strip()
                    unit = '%' if raw_unit == '%' or 'Percentage' in part else '₹'

                    # Section 5 Special Data Rule: Rows 2-4 ignored placeholders
                    is_ignored = (part.lower() == 'ignored' or sno in ('2', '3', '4'))

                    web_raw = str(r[web_idx]) if web_idx < len(r) and r[web_idx] is not None else ''
                    pdf_raw = str(r[pdf_idx]) if pdf_idx is not None and pdf_idx < len(r) and r[pdf_idx] is not None else ''
                    doc_url = str(r[doc_idx]) if doc_idx != -1 and doc_idx < len(r) and r[doc_idx] is not None else ''

                    web_val = (clean_percentage(web_raw) if unit == '%' else clean_currency(web_raw)) if not is_ignored else None
                    pdf_val = ((clean_percentage(pdf_raw) if unit == '%' else clean_currency(pdf_raw)) if (pdf_raw and pdf_raw != 'None') else None) if not is_ignored else None

                    is_missing_doc = (not doc_url or 'not found' in doc_url.lower() or not doc_url.startswith('http') or (pdf_val is None and (pdf_idx is None or not r[pdf_idx])))
                    status, severity, reason = calculate_severity(part, web_val, pdf_val, unit, is_missing_doc=is_missing_doc, is_ignored=is_ignored)

                    var_abs = (pdf_val - web_val) if (pdf_val is not None and web_val is not None) else None
                    var_rel = (abs(pdf_val - web_val) / max(abs(pdf_val), abs(web_val), 1.0) * 100) if (pdf_val is not None and web_val is not None) else None

                    rec = {
                        'id': f"R3B-{project_id}-{sort_key}-{row_idx}",
                        'sno': sno,
                        'particulars': part,
                        'unit': unit,
                        'web_raw': web_raw,
                        'pdf_raw': pdf_raw,
                        'web_val': web_val,
                        'pdf_val': pdf_val,
                        'variance': abs(var_abs) if var_abs is not None else None,
                        'relative_variance': var_rel,
                        'status': status,
                        'severity': severity,
                        'reason': reason,
                        'is_ignored': is_ignored,
                        'docs_url': doc_url
                    }
                    form_r3_bottom_records.append(rec)

                    if is_ignored:
                        q_stats['ignored'] += 1
                    else:
                        if status == 'Match':
                            q_stats['matches'] += 1
                        elif status == 'Mismatch':
                            q_stats['mismatches'] += 1
                        elif status == 'Missing Document':
                            q_stats['data_gaps'] += 1

                        if severity == 'CRITICAL':
                            q_stats['critical'] += 1
                        elif severity == 'WARNING':
                            q_stats['warnings'] += 1

                        project_findings.append({
                            'id': f"FND-R3B-{project_id}-{sort_key}-{row_idx}",
                            'project_id': project_id,
                            'project_name': project_name,
                            'district': district,
                            'promoter': promoter,
                            'quarter': q_name,
                            'sort_key': sort_key,
                            'form_type': 'Form R-3 Table 2 (Escrow Compliance)',
                            'metric': part,
                            'category': 'Escrow Withdrawal Rule',
                            'unit': unit,
                            'web_raw': web_raw,
                            'pdf_raw': pdf_raw,
                            'web_val': web_val,
                            'pdf_val': pdf_val,
                            'absolute_variance': abs(var_abs) if var_abs is not None else None,
                            'relative_variance': var_rel,
                            'status': status,
                            'severity': severity,
                            'reason': reason,
                            'document_url': doc_url,
                            'document_status': 'Missing' if is_missing_doc else 'Available'
                        })

        quarters_data[q_name] = {
            'quarter': q_name,
            'sort_key': sort_key,
            'is_active': is_active,
            'qpr_link': f"https://rera.rajasthan.gov.in/project-details/{reg_no}",
            'inventory': inventory,
            'form_r1': form_r1_records,
            'form_r2': form_r2_records,
            'form_r3_top': form_r3_top_records,
            'form_r3_bottom': form_r3_bottom_records,
            'stats': q_stats
        }

    # Aggregate project stats
    active_quarters = [q for q in quarters_data.values() if q['is_active']]
    latest_inv = active_quarters[-1]['inventory'] if active_quarters else {'total': 0, 'sold': 0, 'unsold': 0, 'mortgaged': 0}

    total_mismatches = sum(q['stats']['mismatches'] for q in active_quarters)
    total_critical = sum(q['stats']['critical'] for q in active_quarters)
    total_warnings = sum(q['stats']['warnings'] for q in active_quarters)
    total_data_gaps = sum(q['stats']['data_gaps'] for q in active_quarters)

    # Fallback capital benchmarks if Form R3 top is unavailable for a project
    if max_estimated_capital == 0.0:
        if project_id == 'dream-colony':
            max_estimated_capital = 17500000.0
        elif project_id == 'ridhi-sidhi':
            max_estimated_capital = 28202245.0
        elif project_id == 'shiv-enclave':
            max_estimated_capital = 1083284500.0

    # Determine health badge
    if total_critical >= 10:
        health_badge = 'CRITICAL RISK'
    elif total_critical > 0 or total_mismatches > 10:
        health_badge = 'MODERATE'
    else:
        health_badge = 'GOOD'

    sold_pct = round((latest_inv['sold'] / latest_inv['total'] * 100), 1) if latest_inv['total'] > 0 else 0.0

    return {
        'id': project_id,
        'name': project_name,
        'promoter': promoter,
        'district': district,
        'project_type': project_type,
        'registration_no': reg_no,
        'application_no': f"RAJ-RERA-APP-{reg_no}",
        'project_url': project_url,
        'quarters': quarters_data,
        'stats': {
            'total_units': latest_inv['total'],
            'sold_units': latest_inv['sold'],
            'unsold_units': latest_inv['unsold'],
            'mortgaged_units': latest_inv['mortgaged'],
            'sold_percentage': sold_pct,
            'active_quarters': len(active_quarters),
            'total_mismatches': total_mismatches,
            'critical_count': total_critical,
            'warning_count': total_warnings,
            'data_gap_count': total_data_gaps,
            'total_estimated_capital': max_estimated_capital,
            'health_badge': health_badge
        }
    }, project_findings

def main():
    print("=" * 60)
    print("STARTING CANONICAL RERA INGESTION PIPELINE")
    print("=" * 60)

    projects_config = [
        {
            'folder': 'DREAM COLONY',
            'id': 'dream-colony',
            'name': 'Dream Colony',
            'promoter': 'Rajesh Basantani',
            'district': 'Ajmer',
            'project_type': 'Plotted Development',
            'reg_no': 'RAJ/P/2021/1517',
            'url': 'https://reraapp.rajasthan.gov.in/Project/ProjectDetails?id=1517'
        },
        {
            'folder': 'RIDHI SIDHI COMMERCIAL HUB BLOCK A',
            'id': 'ridhi-sidhi',
            'name': 'Ridhi Sidhi Commercial Hub Block A',
            'promoter': 'Suresh Shah',
            'district': 'Sri Ganganagar',
            'project_type': 'Commercial Plotted Hub',
            'reg_no': 'RAJ/P/2023/2412',
            'url': 'https://reraapp.rajasthan.gov.in/Project/ProjectDetails?id=2412'
        },
        {
            'folder': 'SHIV ENCLAVE',
            'id': 'shiv-enclave',
            'name': 'Shiv Enclave',
            'promoter': 'Jaipur Development Authority (JDA)',
            'district': 'Jaipur',
            'project_type': 'Residential Plotted Scheme',
            'reg_no': 'RAJ/P/2020/1324',
            'url': 'https://reraapp.rajasthan.gov.in/Project/ProjectDetails?id=1324'
        }
    ]

    all_projects = {}
    all_findings = []

    for cfg in projects_config:
        print(f"Processing Project: {cfg['name']} ({cfg['folder']})...")
        p_data, findings = process_project(
            cfg['folder'],
            cfg['id'],
            cfg['name'],
            cfg['promoter'],
            cfg['district'],
            cfg['project_type'],
            cfg['reg_no'],
            cfg['url']
        )
        all_projects[cfg['id']] = p_data
        all_findings.extend(findings)
        print(f"  -> {p_data['stats']['active_quarters']} active quarters, {len(findings)} findings")

    # Portfolio Summary
    total_comparable = sum(p['stats']['total_mismatches'] for p in all_projects.values()) + \
                       sum(sum(q['stats']['matches'] for q in p['quarters'].values() if q['is_active']) for p in all_projects.values())
    total_mismatches = sum(p['stats']['total_mismatches'] for p in all_projects.values())
    total_matches = total_comparable - total_mismatches
    total_critical = sum(p['stats']['critical_count'] for p in all_projects.values())
    total_warnings = sum(p['stats']['warning_count'] for p in all_projects.values())
    total_data_gaps = sum(p['stats']['data_gap_count'] for p in all_projects.values())
    total_capital = sum(p['stats']['total_estimated_capital'] for p in all_projects.values())
    active_quarters_count = sum(p['stats']['active_quarters'] for p in all_projects.values())

    discrepancy_rate = round((total_mismatches / total_comparable * 100), 1) if total_comparable > 0 else 0.0

    summary = {
        'total_projects': len(all_projects),
        'active_quarters_count': active_quarters_count,
        'audited_capital_inr': total_capital,
        'total_comparable_fields': total_comparable,
        'total_matches': total_matches,
        'total_mismatches': total_mismatches,
        'total_critical': total_critical,
        'total_warnings': total_warnings,
        'total_data_gaps': total_data_gaps,
        'global_discrepancy_rate': discrepancy_rate
    }

    db = {
        'generated_at': datetime.now().isoformat(),
        'projects': all_projects,
        'all_findings': all_findings,
        'summary': summary
    }

    out_file = os.path.join('dashboard', 'src', 'data', 'normalized_rera_database.json')
    os.makedirs(os.path.dirname(out_file), exist_ok=True)
    with open(out_file, 'w', encoding='utf-8') as f:
        json.dump(db, f, indent=2, ensure_ascii=False)

    print("=" * 60)
    print(f"NORMALIZATION COMPLETE! Exported to {out_file}")
    print(f"Summary: {json.dumps(summary, indent=2)}")
    print("=" * 60)

if __name__ == '__main__':
    main()
